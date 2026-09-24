import pytest

from dictify import exporters
from dictify.formatting import MODE_SEGMENTS, ViewOptions


@pytest.mark.parametrize("fmt", [f.key for f in exporters.FORMATS])
@pytest.mark.parametrize("mode", ["transcript", MODE_SEGMENTS])
def test_export_all_formats(qapp, transcript, tmp_path, fmt, mode):
    path = tmp_path / f"out.{fmt}"
    exporters.export(transcript, ViewOptions(mode=mode), fmt, path)
    assert path.stat().st_size > 0


def test_docx_content(transcript, tmp_path):
    from docx import Document

    path = tmp_path / "out.docx"
    exporters.export(transcript, ViewOptions(), "docx", path)
    text = [p.text for p in Document(str(path)).paragraphs]
    assert text == ["interview", "Hello and welcome. Tell me about yourself.", "I love working with people. I was a teacher."]

    exporters.export(transcript, ViewOptions(mode=MODE_SEGMENTS, end_times=False), "docx", path)
    text = [p.text for p in Document(str(path)).paragraphs]
    assert text[1] == "[00:00]  Hello and welcome."


def test_pdf_is_valid(qapp, transcript, tmp_path):
    path = tmp_path / "out.pdf"
    transcript.segments[0].text = "Привет, мир."
    exporters.export(transcript, ViewOptions(), "pdf", path)
    assert path.read_bytes().startswith(b"%PDF")


def test_unknown_format(transcript, tmp_path):
    with pytest.raises(ValueError):
        exporters.export(transcript, ViewOptions(), "xyz", tmp_path / "x")
