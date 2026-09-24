"""Writes a transcript to disk in the supported export formats."""
from __future__ import annotations

import html
from dataclasses import dataclass
from pathlib import Path

from dictify.formatting import (
    MODE_SEGMENTS,
    ViewOptions,
    block_text,
    blocks,
    to_markdown,
    to_srt,
    to_text,
)
from dictify.model import Transcript

MUTED_RGB = (0x8A, 0x8A, 0x8E)


@dataclass(frozen=True)
class ExportFormat:
    key: str
    label: str
    ext: str


FORMATS = [
    ExportFormat("txt", "Plain text (.txt)", ".txt"),
    ExportFormat("md", "Markdown (.md)", ".md"),
    ExportFormat("srt", "Subtitles (.srt)", ".srt"),
    ExportFormat("docx", "Word (.docx)", ".docx"),
    ExportFormat("pdf", "PDF (.pdf)", ".pdf"),
]


def export(t: Transcript, opts: ViewOptions, fmt: str, path: str | Path) -> None:
    path = Path(path)
    if fmt == "txt":
        path.write_text(to_text(t, opts), encoding="utf-8")
    elif fmt == "md":
        path.write_text(to_markdown(t, opts), encoding="utf-8")
    elif fmt == "srt":
        path.write_text(to_srt(t), encoding="utf-8")
    elif fmt == "docx":
        _write_docx(t, opts, path)
    elif fmt == "pdf":
        _write_pdf(t, opts, path)
    else:
        raise ValueError(f"Unknown export format: {fmt}")


def _write_docx(t: Transcript, opts: ViewOptions, path: Path) -> None:
    from docx import Document
    from docx.shared import Pt, RGBColor

    title = Path(t.media_path).stem
    doc = Document()
    doc.core_properties.title = title
    doc.add_heading(title, level=1)

    def gray(run):
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(*MUTED_RGB)

    for b in blocks(t, opts):
        p = doc.add_paragraph()
        if b.label:
            gray(p.add_run(f"[{b.label}]  "))
        p.add_run(block_text(t, b))
    doc.save(str(path))


def _write_pdf(t: Transcript, opts: ViewOptions, path: Path) -> None:
    from PySide6.QtCore import QMarginsF
    from PySide6.QtGui import QFont, QPageLayout, QPageSize, QPdfWriter, QTextDocument

    title = Path(t.media_path).stem
    writer = QPdfWriter(str(path))
    writer.setTitle(title)
    writer.setCreator("Dictify")
    writer.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
    # QTextDocument.print_ adds its own 2 cm margin around the text.
    writer.setPageMargins(QMarginsF(0, 0, 0, 0), QPageLayout.Unit.Millimeter)

    doc = QTextDocument()
    font = QFont()
    font.setPointSize(11)
    doc.setDefaultFont(font)
    doc.setHtml(transcript_html(t, opts))
    doc.print_(writer)


def transcript_html(t: Transcript, opts: ViewOptions) -> str:
    muted = "#%02x%02x%02x" % MUTED_RGB
    parts = [f"<h2>{html.escape(Path(t.media_path).stem)}</h2>"]
    for b in blocks(t, opts):
        text = html.escape(block_text(t, b))
        if opts.mode == MODE_SEGMENTS:
            label = f'<span style="color:{muted}">[{html.escape(b.label)}]</span>&nbsp; ' if b.label else ""
            parts.append(f'<p style="margin:0 0 6px 0">{label}{text}</p>')
        else:
            if b.label:
                parts.append(
                    f'<p style="color:{muted};font-size:9pt;margin:10px 0 2px 0">{html.escape(b.label)}</p>'
                )
            parts.append(f'<p style="margin:0 0 8px 0">{text}</p>')
    return "\n".join(parts)
